import os
import re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment


def _sanitize_filename_part(value: str) -> str:
    sanitized = re.sub(r'[<>:"/\\|?*]+', "_", value).strip().rstrip(".")
    return sanitized or "report"

def create_excel_report(class_name, subject_name, grades_data):
    """
    Creates an Excel report and returns the file path.
    grades_data: list of dicts {'full_name': str, 'grades': list[int]}
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Ведомость оценок"

    # Header
    ws['A1'] = f"Ведомость оценок: {class_name}"
    ws['A2'] = f"Предмет: {subject_name}"
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'].font = Font(bold=True)
    
    ws.append([]) # Empty row
    
    # Table headers
    headers = ["№", "ФИО ученика", "Оценки", "Средний балл"]
    ws.append(headers)
    
    header_font = Font(bold=True)
    for cell in ws[4]:
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Data
    for i, item in enumerate(grades_data, start=1):
        name = item['full_name']
        grades = item['grades']
        grades_str = ", ".join(map(str, grades)) if grades else ""
        avg = round(sum(grades) / len(grades), 2) if grades else 0.0
        
        ws.append([i, name, grades_str, avg])

    # Styling columns
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 15
    ws.freeze_panes = "A5"

    # Save to temp file
    os.makedirs("temp_reports", exist_ok=True)
    safe_class_name = _sanitize_filename_part(class_name).replace(" ", "_")
    safe_subject_name = _sanitize_filename_part(subject_name).replace(" ", "_")
    filename = f"report_{safe_class_name}_{safe_subject_name}.xlsx"
    file_path = os.path.join("temp_reports", filename)
    wb.save(file_path)
    
    return file_path
